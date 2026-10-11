"""Build the Channel Performance summary workbook."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from pytool.constants import DEFAULT_CATEGORY_HEADER

GREEN_FONT = Font(color="006100")
RED_FONT = Font(color="9C0006")
GREEN_FILL = PatternFill("solid", fgColor="C6EFCE")
RED_FILL = PatternFill("solid", fgColor="FFC7CE")

QUAD_KEYS = ("direct_ic", "direct_ooc", "non_direct_ic", "non_direct_ooc")
OTHER_BLANK = (
    {"name": "Other", "key": "other"},
    {"name": "Blank", "key": "blank"},
)


def _unique_bucket_count(rows: list, prospect_id_index: int) -> int:
    if prospect_id_index < 0:
        return len(rows)
    seen: set[str] = set()
    blank_count = 0
    for row in rows:
        if prospect_id_index >= len(row) or row[prospect_id_index] is None:
            blank_count += 1
            continue
        prospect_id = str(row[prospect_id_index]).strip()
        if prospect_id:
            seen.add(prospect_id)
        else:
            blank_count += 1
    return len(seen) + blank_count


def _source_columns(lead_categories) -> list:
    selected = [
        cat for cat in (lead_categories or [])
        if cat.get("key") not in {"other", "blank"}
    ]
    return selected + list(OTHER_BLANK)


def month_counts(
    new_this_month: int,
    total_mailable: int,
    buckets: dict,
    lead_categories=None,
    prospect_id_index: int = -1,
) -> list:
    def bucket_count(key: str) -> int:
        return _unique_bucket_count(buckets[key], prospect_id_index)

    return [
        total_mailable,
        new_this_month,
        bucket_count("direct_ic"),
        bucket_count("direct_ooc"),
        bucket_count("non_direct_ic"),
        bucket_count("non_direct_ooc"),
        *[bucket_count(cat["key"]) for cat in (lead_categories or [])],
        bucket_count("other"),
        bucket_count("blank"),
    ]


def _mom(current, prior):
    current = current or 0
    prior = prior or 0
    if prior == 0:
        return 0 if current == 0 else None
    return (current - prior) / prior


def write_channel_performance(
    path: Path,
    current_month: str,
    # original rows
    current_rows: list,
    # filtered_buckets
    current_buckets: dict,
    prior_month: str | None,
    prior_rows: list,
    prior_buckets: dict,
    current_mailable: int,
    prior_mailable: int,
    new_this_month: int,
    new_prior_month: int,
    prospect_id_index: int = -1,
    headers=None,
    lead_categories=None,
) -> Path:
    sources = _source_columns(lead_categories)
    current_vals = month_counts(
        new_this_month, current_mailable, current_buckets, sources,
        prospect_id_index)
    prior_vals = month_counts(
        new_prior_month, prior_mailable, prior_buckets, sources,
        prospect_id_index)

    wb = Workbook()
    ws = wb.active
    ws.title = "Channel Performance"
    ws.append(["Month"] + [c["name"] for c in DEFAULT_CATEGORY_HEADER + sources])
    ws.append([prior_month or ""] + prior_vals)
    ws.append([current_month] + current_vals)

    #   pre  [ 10, 12 ]
    #   curr  [ 11, 15 ]
    # zip(10, 11), zip(12, 15)
    mom_vals = [_mom(c, p) for c, p in zip(current_vals, prior_vals)]
    ws.append(["MoM % Change"] + mom_vals)
    # tint the value based on whether it is positive or negative
    for col, val in enumerate(mom_vals, start=2):
        cell = ws.cell(row=4, column=col)
        cell.number_format = "0.0%"
        if val is not None and val > 0:
            cell.font = GREEN_FONT
            cell.fill = GREEN_FILL
        elif val is not None and val < 0:
            cell.font = RED_FONT
            cell.fill = RED_FILL

    wb.save(path)
    return path
