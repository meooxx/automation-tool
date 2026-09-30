"""Fill Master report sample.xlsx with 100 fully populated source rows."""

from __future__ import annotations

import random
from datetime import date, datetime, time, timedelta
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / "templates" / "Master report sample.xlsx"
OUT = ROOT / "data" / "master_source_100.xlsx"

LEAD_SOURCES = [
    "Career Match",
    "Connect W Recr",
    "Recruitment events",
    "JOML",
    "Still time to apply",
    "Other",
    "FOH",
    "AppDay",
    "Auto Web Form",
    "SOH",
    "Tell Us More",
    "GC Website",
    "int recruitment",
]
RESD_IC = ["IC", "ICC", "in-catchment", "in catchment"]
RESD_OOC = ["OOC", "OCC", "out-of-catchment", "out of catchment"]
APP_DIRECT = [
    "Direct from Secondary School",
    "Direct from High School",
    "Direct",
]
APP_NON_DIRECT = [
    "Non-Direct",
    "Non Direct",
    "nondirect",
    "Mature Student",
    "College Transfer",
    "University Transfer",
]
QUAD_VALUES = [
    (resd, app)
    for resd in (*RESD_IC, *RESD_OOC)
    for app in (*APP_DIRECT, *APP_NON_DIRECT)
]
FIRST = [
    "Emma", "Liam", "Noah", "Olivia", "Sophia", "Lucas", "Mason", "Mia",
    "James", "Ella", "Ava", "Ethan", "Isabella", "Logan", "Harper", "Jack",
    "Lily", "Henry", "Grace", "Owen", "Chloe", "Leo", "Zoe", "Nathan",
    "Aria", "Caleb", "Nora", "Ryan", "Layla", "Isaac",
]
LAST = [
    "Wei", "Park", "Singh", "Chen", "Patel", "Brown", "Nguyen", "Garcia",
    "Kim", "Rossi", "Ahmed", "Martin", "Lopez", "Khan", "Wilson",
]
CITIES = [
    ("Barrie", "ON", "L4N 1A1"),
    ("Orillia", "ON", "L3V 2B2"),
    ("Midland", "ON", "L4R 3C3"),
    ("Toronto", "ON", "M5V 1J1"),
    ("Ottawa", "ON", "K1A 0B1"),
    ("Mississauga", "ON", "L5B 2C4"),
]
CAMPUSES = ["Barrie", "Orillia", "Midland", "Owen Sound"]
PROGRAMS = [
    ("Computer Programming", "CMPG", "Technology"),
    ("Business", "BUSN", "Business"),
    ("Nursing", "NURS", "Health"),
    ("Automotive", "AUTO", "Skilled Trades"),
    ("Hospitality", "HOSP", "Hospitality"),
    ("Police Foundations", "PFND", "Community Safety"),
]
HIGH_SCHOOLS = ["Barrie Central", "Orillia Secondary", "Eastview", "St. Joseph", "Newmarket HS"]
SALUTATIONS = ["Ms.", "Mr.", "Mx."]
GRADES = ["A", "B", "C"]
SHIRTS = ["XS", "S", "M", "L", "XL"]
EQUITY = ["Prefer not to say", "Visible minority", "Person with a disability", "None identified"]
STATUS_CA = ["Citizen", "Permanent Resident", "Study Permit", "Other"]
DELIVERY = ["In person", "Online", "Hybrid"]
ACTIVITY = ["Active", "Inactive", "New"]
AGENTS = ["North Star Agency", "Campus Link", "Direct"]
LISTS = ["Mailable", "Viewbook 2026", "Event Follow-up"]
YES_NO = ("Yes", "No")
SPEC = [
    ("P101", "Emma", "Wei", "emma@test.com", date(2026, 9, 2), "Career Match", "IC", "Direct from Secondary School"),
    ("P102", "Liam", "Park", "liam@test.com", date(2026, 9, 5), "FOH", "ICC", "Mature Student"),
    ("P103", "Noah", "Singh", "noah@test.com", date(2026, 9, 12), "AppDay", "OOC", "Direct from High School"),
    ("P104", "Olivia", "Chen", "olivia@test.com", date(2026, 9, 8), "Connect W Recr", "in catchment", "College Transfer"),
    ("P105", "Sophia", "Patel", "sophia@test.com", date(2026, 9, 14), "Auto Web Form", "OCC", "Direct"),
    ("P106", "Lucas", "Brown", "lucas@test.com", date(2026, 9, 10), "Career Match", "out-of-catchment", "Non-Direct"),
    ("P107", "Mason", "Nguyen", "mason@test.com", date(2026, 9, 19), "GC Website", "in-catchment", "University Transfer"),
    ("P108", "Mia", "Garcia", "mia@test.com", date(2026, 9, 22), "Other", "out of catchment", "nondirect"),
    ("P099", "James", "Kim", "james@test.com", date(2026, 8, 15), "Career Match", "IC", "Direct from Secondary School"),
    ("P098", "Ella", "Rossi", "ella@test.com", date(2026, 8, 20), "FOH", "OOC", "Non Direct"),
]


def joined_for(i: int) -> date:
    if i < 70:
        return date(2026, 9, 1) + timedelta(days=i % 29)
    if i < 95:
        return date(2026, 8, 1) + timedelta(days=(i - 70) % 30)
    return date(2026, 7, 3) + timedelta(days=(i - 95) % 20)


def build_people(n: int = 100) -> list[dict]:
    rng = random.Random(26)
    people: list[dict] = []
    for item in SPEC:
        pid, first, last, email, joined, source, catchment, applicant_type = item
        people.append(
            {
                "pid": pid,
                "first": first,
                "last": last,
                "email": email,
                "joined": joined,
                "source": source,
                "catchment": catchment,
                "applicant_type": applicant_type,
                "opt_out": False,
            }
        )
    used = {p["pid"] for p in people}
    seq = 1
    while len(people) < n:
        pid = f"P{seq:03d}"
        seq += 1
        if pid in used:
            continue
        first = FIRST[len(people) % len(FIRST)]
        last = LAST[len(people) % len(LAST)]
        catchment, applicant_type = QUAD_VALUES[len(people) % len(QUAD_VALUES)]
        people.append(
            {
                "pid": pid,
                "first": first,
                "last": last,
                "email": f"{first.lower()}.{last.lower()}.{seq}@test.com",
                "joined": joined_for(len(people)),
                "source": LEAD_SOURCES[len(people) % len(LEAD_SOURCES)],
                "catchment": catchment,
                "applicant_type": applicant_type,
                "opt_out": rng.random() < 0.06,
            }
        )
        used.add(pid)
    return people


def yn(flag: bool) -> str:
    return "Yes" if flag else "No"


def dt(day: date, hour: int = 9, minute: int = 0) -> datetime:
    return datetime.combine(day, time(hour, minute))


def value_for(header: str, p: dict, i: int, rng: random.Random):
    city, state, zip_code = CITIES[i % len(CITIES)]
    campus = CAMPUSES[i % len(CAMPUSES)]
    program, code, area = PROGRAMS[i % len(PROGRAMS)]
    program2, code2, area2 = PROGRAMS[(i + 1) % len(PROGRAMS)]
    program3, code3, area3 = PROGRAMS[(i + 2) % len(PROGRAMS)]
    joined = p["joined"]
    created = joined - timedelta(days=rng.randint(0, 8))
    updated = joined + timedelta(days=rng.randint(0, 5))
    activity = datetime.combine(joined, time(14, 30)) + timedelta(hours=rng.randint(1, 48))
    personal = p["email"].replace("@test.com", ".personal@gmail.com")
    college_email = f"{p['first'][0].lower()}{p['last'].lower()}@georgiancollege.ca"
    owner = f"recruiter{(i % 8) + 1}@georgiancollege.ca"
    advisor = f"advisor{(i % 5) + 1}@georgiancollege.ca"
    hs = HIGH_SCHOOLS[i % len(HIGH_SCHOOLS)]
    yob = 2005 + (i % 6)
    georgian_id = f"G{2000000 + i}"
    ocas = f"OC{300000 + i}"
    street = f"{100 + i} College Dr"
    opt = p["opt_out"]
    source = p["source"] or "Other"
    catchment = p["catchment"]
    medium = ["organic", "paid", "email", "event", "referral"][i % 5]
    campaign = f"Fall-2026-{source.replace(' ', '')}"
    shirt = SHIRTS[i % len(SHIRTS)]
    guests = (i % 4)
    convocation = date(2026, 6, 12)
    test_day = date(2026, 10, 6)
    train1 = date(2026, 10, 14)
    train2 = date(2026, 10, 21)
    nii = date(2026, 11, 3)

    by_header = {
        "Prospect Id": p["pid"],
        "First Name": p["first"],
        "Last Name": p["last"],
        "Email": p["email"],
        "Company": "N/A",
        "Last Activity At": activity,
        "Campaign": campaign,
        "Notes": f"Imported test prospect {p['pid']}",
        "Score": 40 + (i % 60),
        "Grade": GRADES[i % 3],
        "Website": "https://www.georgiancollege.ca",
        "Job Title": "Student",
        "Department": area,
        "Country": "Canada",
        "Address One": street,
        "Address Two": f"Unit {i}",
        "City": city,
        "State": state,
        "Territory": "Ontario",
        "Zip": zip_code,
        "Phone": f"705-555-{1000 + i:04d}",
        "Fax": f"705-554-{1000 + i:04d}",
        "Source": source,
        "Annual Revenue": 0,
        "Employees": 0,
        "Industry": "Education",
        "Do Not Email": yn(opt),
        "Do Not Call": yn(opt),
        "Years In Business": 0,
        "Comments": "Synthetic master-report row for automation testing",
        "Salutation": SALUTATIONS[i % 3],
        "Opted Out": yn(opt),
        "Referrer": LAST[(i + 3) % len(LAST)] + " High School",
        "Created Date": created,
        "Updated Date": updated,
        "User": owner,
        "First Assigned": owner,
        "CRM Contact FID": f"003{i:011d}",
        "CRM Lead FID": f"00Q{i:011d}",
        "Joined": joined,
        "Opted Out of List": yn(opt),
        "Lists": LISTS[i % len(LISTS)],
        "A A1 Prog Code": code,
        "A A2 Prog Code": code2,
        "A A3 Prog Code": code3,
        "A OCAS ID": ocas,
        "A OCAS Status": "Applied",
        "A OCAS Year": 2026,
        "A P1 Prog Code": code,
        "A P2 Prog Code": code2,
        "A P3 Prog Code": code3,
        "Access and Pathway Advisor Email": advisor,
        "Accessibility needs": "None" if i % 7 else "Exam accommodations",
        "Active Program Enrollment Registered": yn(i % 9 == 0),
        "Activity Status of Contact": ACTIVITY[i % 3],
        "Agent Name": AGENTS[i % len(AGENTS)],
        "Agent Status": "Active",
        "B A1 Prog Code": code2,
        "B A2 Prog Code": code3,
        "B A3 Prog Code": code,
        "B OCAS ID": f"B{ocas}",
        "B OCAS Status": "Not applied",
        "B OCAS Year": 2027,
        "B P1 Prog Code": code2,
        "B P2 Prog Code": code3,
        "B P3 Prog Code": code,
        "Campus of Interest": campus,
        "Campus of Interest-OH": campus,
        "Classification Prospect": catchment,
        "College Email": college_email,
        "College Entry Advisor": advisor,
        "College Name ": "Georgian College",
        "Congrats Video ID": f"VID-{i:04d}",
        "Contact Description": "Prospective student",
        "Contact Description - Influencer": "Guidance counsellor" if i % 11 == 0 else "Self",
        "Contact Owner Email": owner,
        "Contact Record Type": "Prospect",
        "Content": f"{source} landing page",
        "Contest and prize information": yn(i % 13 == 0),
        "Convocation Date": convocation,
        "Convocation Gown Time": time(8, 30),
        "Convocation Program Title": program,
        "Convocation Start Time": time(10, 0),
        "Deceased": "No",
        "Dietary restrictions or allergies": "None" if i % 8 else "Vegetarian",
        "Do Not Contact": yn(opt),
        "Email Bounced": "No",
        "Email Opt In": yn(not opt),
        "Equity-deserving groups": EQUITY[i % len(EQUITY)],
        "Event delivery": DELIVERY[i % 3],
        "Event Information": f"{source} session",
        "Event Notes": "Registered",
        "Faculty Status": "N/A",
        "Georgian ID": georgian_id,
        "Graduation year": yob + 18,
        "Guests": guests,
        "Guidance & Student Success Status": "Open",
        "High School Name": hs,
        "How can we help?": "Program information",
        "ICR Regional Email": f"icr.{campus.lower().replace(' ', '')}@georgiancollege.ca",
        "Indigenous Self-identification": "No" if i % 10 else "Yes",
        "ISV School Name": hs,
        "last form completed date": joined,
        "Last Modified Date": updated,
        "Lead Source - Most Recent": source,
        "Mailing Address": f"{street}, {city}, {state} {zip_code}",
        "MCR Contact ID": f"MCR-{p['pid']}",
        "Medium": medium,
        "Newcomer": yn(i % 6 == 0),
        "Nomination for GNED 1090": yn(i % 14 == 0),
        "Nominee name (if nominating someone else)": f"{FIRST[(i + 4) % len(FIRST)]} {LAST[(i + 4) % len(LAST)]}",
        "OH Applicant Type": p["applicant_type"],
        "Permissions Method": "Web form",
        "Personal Email": personal,
        "Phone Extension": str(200 + i),
        "Point of Entry": "Semester 1",
        "Preferred Email": p["email"],
        "Preferred First Name": p["first"],
        "Preferred Name": p["first"],
        "Primary Academic Area of Interest": area,
        "Primary Program of Interest": program,
        "Program Campus": campus,
        "Program Delivery": DELIVERY[i % 3],
        "Program/Course of interest": program,
        "Project Management at the NII Dates": nii,
        "Reason for Nomination": "Academic achievement" if i % 14 == 0 else "N/A",
        "Record Type": "Prospect",
        "Resd Code IRx": catchment,
        "Secondary Academic Area of Interest": area2,
        "Secondary Program of Interest": program2,
        "Secondary School": hs,
        "Services of Interest": "Campus tour; Advising",
        "Shirt size": shirt,
        "Status in Canada": STATUS_CA[i % len(STATUS_CA)],
        "Status In Canada - MKTG": STATUS_CA[i % len(STATUS_CA)],
        "Tertiary Academic Area of Interest": area3,
        "Tertiary Program of Interest": program3,
        "Test Date Offered": test_day,
        "Training session dates": train1,
        "Training session dates2": train2,
        "Type of Activity": "Inquiry",
        "Updates for future students": yn(not opt),
        "Updates for teachers and counsellors": yn(i % 5 == 0),
        "Viewbook URLs": "https://www.georgiancollege.ca/viewbook",
        "Visa Recieved": yn(STATUS_CA[i % len(STATUS_CA)] == "Study Permit"),
        "Year of Birth": yob,
        "GA Campaign": campaign,
        "GA Medium": medium,
        "GA Source": source,
        "GA Content": f"{source}-hero",
        "GA Term": "fall-intake",
        "Score - Google Ad": 10 + (i % 90),
    }
    if header not in by_header:
        raise KeyError(header)
    return by_header[header]


def write_cell(ws, row: int, col: int, value) -> None:
    cell = ws.cell(row=row, column=col, value=value)
    if isinstance(value, datetime):
        cell.number_format = "YYYY-MM-DD HH:MM"
    elif isinstance(value, date):
        cell.number_format = "YYYY-MM-DD"
    elif isinstance(value, time):
        cell.number_format = "HH:MM"


def main() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(TEMPLATE)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    if any(h is None or str(h).strip() == "" for h in headers):
        raise SystemExit("template has empty header cells")

    rng = random.Random(26)
    people = build_people(100)
    for row_idx, person in enumerate(people, start=2):
        for col, header in enumerate(headers, start=1):
            write_cell(ws, row_idx, col, value_for(str(header), person, row_idx - 1, rng))

    last_col = get_column_letter(len(headers))
    ws.auto_filter.ref = f"A1:{last_col}{1 + len(people)}"
    wb.save(OUT)

    sample = people[0]
    empty = 0
    for col in range(1, len(headers) + 1):
        if ws.cell(2, col).value in (None, ""):
            empty += 1
    print(f"wrote {len(people)} rows x {len(headers)} columns -> {OUT}")
    print(f"empty cells on row 2: {empty}")
    print(f"sample {sample['pid']} {sample['first']} {sample['source']} {sample['joined']}")


if __name__ == "__main__":
    main()
